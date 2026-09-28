from __future__ import annotations

from collections import Counter
import copy
import json
from pathlib import Path
import re
import shutil
from typing import Any

from .table_processing import _notes, _render_candidate_preview, _write_workbook, _write_json
from materials_workbench.storage import candidates as load_candidates, write_index
from materials_workbench.storage import source_page
import pymupdf
from .identifiers import parse_identifier_prefix


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _candidate_weight(candidate: dict[str, Any]) -> int:
    pages = len(candidate.get("source_physical_pages", []))
    cells = len(candidate.get("grid", {}).get("cells", []))
    codes = {issue.get("code") for issue in candidate.get("issues", [])}
    weight = 1 + max(0, pages - 1) + (1 if cells > 100 else 0)
    if codes & {
        "cross_fragment_column_mismatch",
        "multiple_html_tables",
        "large_merged_region",
        "duplicate_identifier_candidates",
    }:
        weight += 2
    if codes & {"embedded_media", "catalog_identifier_collision"}:
        weight += 3
    return weight


def plan_luna_batches(
    *,
    run_dir: Path,
    contract: Path,
    force: bool = False,
    target_weight: int = 9,
    max_tables: int = 10,
    max_pages: int = 20,
    max_cells: int = 500,
) -> dict[str, Any]:
    run_dir = run_dir.expanduser().resolve()
    contract = contract.expanduser().resolve()
    if not contract.is_file():
        raise FileNotFoundError(contract)
    contract_bytes = contract.read_bytes()
    queue = run_dir
    candidates: list[dict[str, Any]] = []
    for table_dir in sorted(queue.glob("Table_*"), key=lambda p: p.name.casefold()):
        candidate_path = table_dir / "candidate.json"
        if not candidate_path.is_file():
            continue
        candidate = _read_json(candidate_path)
        if candidate.get("schema_version") != 2:
            raise ValueError("Only workbench candidate schema 2 is supported")
        if candidate["data_status"] != "review_required":
            continue
        pages = [int(value) for value in candidate.get("source_physical_pages", [])]
        cells = len(candidate.get("grid", {}).get("cells", []))
        candidates.append(
            {
                "candidate_id": candidate["candidate_id"],
                "base_revision": candidate["revision"],
                "table_directory": str(table_dir.resolve()),
                "source_physical_pages": pages,
                "cell_count": cells,
                "weight": _candidate_weight(candidate),
            }
        )

    batches: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    totals = {"weight": 0, "pages": 0, "cells": 0}
    for item in candidates:
        addition = {
            "weight": item["weight"],
            "pages": len(item["source_physical_pages"]),
            "cells": item["cell_count"],
        }
        would_exceed = current and (
            len(current) >= max_tables
            or totals["pages"] + addition["pages"] > max_pages
            or totals["cells"] + addition["cells"] > max_cells
            or totals["weight"] + addition["weight"] > target_weight
        )
        if would_exceed:
            batches.append(current)
            current = []
            totals = {"weight": 0, "pages": 0, "cells": 0}
        current.append(item)
        for key in totals:
            totals[key] += addition[key]
    if current:
        batches.append(current)

    batch_root = run_dir / "luna_batches"
    if batch_root.exists() and force:
        shutil.rmtree(batch_root)
    if batch_root.exists() and any(batch_root.iterdir()):
        raise FileExistsError(f"Luna batch directory is not empty: {batch_root}; use --force")
    batch_root.mkdir(parents=True, exist_ok=True)
    preview_dir = batch_root / "pages"
    preview_dir.mkdir(exist_ok=True)
    page_inputs = {}
    for page in sorted({p for item in candidates for p in item["source_physical_pages"]}):
        source = source_page(run_dir, page)
        image = preview_dir / f"page_{page:04d}.png"
        with pymupdf.open(source) as pdf:
            pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(image)
        page_inputs[page] = {"page": page, "pdf": str(source), "image": str(image.resolve())}
    contract_snapshot = batch_root / "LUNA_TABLE_REVIEW_CONTRACT.md"
    contract_snapshot.write_bytes(contract_bytes)
    records: list[dict[str, Any]] = []
    for index, items in enumerate(batches, 1):
        batch_id = f"batch_{index:03d}"
        batch_dir = batch_root / batch_id
        output_dir = batch_dir / "output"
        (output_dir / "patches").mkdir(parents=True)
        manifest = {
            "schema_version": 1,
            "contract_version": "2.0",
            "batch_id": batch_id,
            "assigned_candidates": [
                {**{key: item[key] for key in ("candidate_id", "base_revision", "table_directory", "source_physical_pages")},
                 "sources": [page_inputs[p] for p in item["source_physical_pages"]]}
                for item in items
            ],
        }
        manifest_path = batch_dir / "input_manifest.json"
        _write_json(manifest_path, manifest)
        records.append(
            {
                "batch_id": batch_id,
                "input_manifest": str(manifest_path.resolve()),
                "output_directory": str(output_dir.resolve()),
                "candidate_ids": [item["candidate_id"] for item in items],
                "totals": {
                    "tables": len(items),
                    "source_pages": sum(len(item["source_physical_pages"]) for item in items),
                    "cells": sum(item["cell_count"] for item in items),
                    "weight": sum(item["weight"] for item in items),
                },
            }
        )
    result = {
        "schema_version": 1,
        "run_dir": str(run_dir),
        "contract_file": str(contract_snapshot.resolve()),
        "review_candidates": len(candidates),
        "batch_count": len(records),
        "batches": records,
    }
    _write_json(run_dir / "luna_batch_plan.json", result)
    return result


def _rebuild_cells(candidate: dict[str, Any]) -> None:
    rows = candidate["grid"]["rows"]
    old = {
        (int(cell["row"]), int(cell["column"])): cell
        for cell in candidate["grid"].get("cells", [])
    }
    cells: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, 1):
        for column_index, value in enumerate(row, 1):
            prior = old.get((row_index, column_index), {})
            cells.append(
                {
                    "cell_id": f"R{row_index:03d}C{column_index:03d}",
                    "row": row_index,
                    "column": column_index,
                    "rowspan": int(prior.get("rowspan", 1)),
                    "colspan": int(prior.get("colspan", 1)),
                    "text": value,
                }
            )
    candidate["grid"]["cells"] = cells


def _cell_position(cell_id: str) -> tuple[int, int]:
    match = re.fullmatch(r"R(\d{3})C(\d{3})", cell_id)
    if not match or int(match.group(1)) == 0 or int(match.group(2)) == 0:
        raise ValueError(f"Invalid cell_id: {cell_id}")
    return int(match.group(1)), int(match.group(2))


def _current_merges(candidate: dict[str, Any]) -> list[str]:
    from openpyxl.utils import get_column_letter

    ranges: list[str] = []
    for cell in candidate["grid"].get("cells", []):
        if int(cell.get("rowspan", 1)) > 1 or int(cell.get("colspan", 1)) > 1:
            row, column = int(cell["row"]), int(cell["column"])
            end_row = row + int(cell.get("rowspan", 1)) - 1
            end_column = column + int(cell.get("colspan", 1)) - 1
            ranges.append(
                f"{get_column_letter(column)}{row}:{get_column_letter(end_column)}{end_row}"
            )
    return sorted(set(ranges))


def _set_merges(candidate: dict[str, Any], ranges: list[str]) -> None:
    from openpyxl.utils.cell import range_boundaries

    for cell in candidate["grid"]["cells"]:
        cell["rowspan"] = 1
        cell["colspan"] = 1
    by_position = {
        (int(cell["row"]), int(cell["column"])): cell
        for cell in candidate["grid"]["cells"]
    }
    for cell_range in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        anchor = by_position.get((min_row, min_col))
        if anchor is None:
            raise ValueError(f"Merge range begins outside the grid: {cell_range}")
        anchor["rowspan"] = max_row - min_row + 1
        anchor["colspan"] = max_col - min_col + 1


def _apply_operation(candidate: dict[str, Any], operation: dict[str, Any]) -> str | None:
    name = operation.get("op")
    target = operation.get("target") or {}
    before = operation.get("before")
    after = operation.get("after")
    rows = candidate["grid"]["rows"]
    note: str | None = None
    if name == "set_cell":
        row, column = _cell_position(str(target.get("cell_id", "")))
        if row > len(rows) or column > len(rows[row - 1]):
            raise ValueError(f"set_cell target is outside grid: {target}")
        if rows[row - 1][column - 1] != before:
            raise ValueError(f"set_cell before mismatch at {target['cell_id']}")
        rows[row - 1][column - 1] = str(after)
    elif name == "insert_row":
        index = int(target["row_index"])
        if before is not None or not isinstance(after, list) or not (1 <= index <= len(rows) + 1):
            raise ValueError(f"Invalid insert_row operation: {operation}")
        rows.insert(index - 1, [str(value) for value in after])
    elif name == "delete_row":
        index = int(target["row_index"])
        if not (1 <= index <= len(rows)) or rows[index - 1] != before or after is not None:
            raise ValueError(f"delete_row before mismatch: {operation}")
        rows.pop(index - 1)
    elif name in {"insert_column", "delete_column"}:
        index = int(target["column_index"])
        width = max((len(row) for row in rows), default=0)
        for row in rows:
            row.extend([""] * (width - len(row)))
        if name == "insert_column":
            if before is not None or not isinstance(after, list) or len(after) != len(rows) or not (1 <= index <= width + 1):
                raise ValueError(f"Invalid insert_column operation: {operation}")
            for row_index, row in enumerate(rows):
                row.insert(index - 1, str(after[row_index]))
        else:
            current = [row[index - 1] for row in rows] if 1 <= index <= width else None
            if current != before or after is not None:
                raise ValueError(f"delete_column before mismatch: {operation}")
            for row in rows:
                row.pop(index - 1)
    elif name in {"merge_cells", "unmerge_cells"}:
        cell_range = str(target.get("range", ""))
        merges = _current_merges(candidate)
        expected = cell_range in merges
        if bool(before) != expected:
            raise ValueError(f"merge before mismatch for {cell_range}")
        if name == "merge_cells":
            if cell_range not in merges:
                merges.append(cell_range)
        elif cell_range in merges:
            merges.remove(cell_range)
        _set_merges(candidate, sorted(merges))
    elif name == "set_metadata":
        field = str(target.get("field", ""))
        if field not in {"caption", "footnotes", "node_id_raw"}:
            raise ValueError(f"Unsupported metadata field: {field}")
        if candidate.get(field) != before:
            raise ValueError(f"set_metadata before mismatch for {field}")
        candidate[field] = str(after)
    elif name == "add_note":
        if target.get("section") != "notes" or before is not None:
            raise ValueError(f"Invalid add_note operation: {operation}")
        note = str(after)
    elif name == "replace_grid":
        if target.get("grid") != "entire" or not isinstance(after, dict) or not isinstance(after.get("rows"), list):
            raise ValueError(f"Invalid replace_grid operation: {operation}")
        summary = {"rows": len(rows), "columns": max((len(row) for row in rows), default=0)}
        if before != summary and before != candidate["grid"]:
            raise ValueError("replace_grid before mismatch")
        candidate["grid"]["rows"] = [[str(value) for value in row] for row in after["rows"]]
        _rebuild_cells(candidate)
        _set_merges(candidate, [str(value) for value in after.get("merges", [])])
    else:
        raise ValueError(f"Unsupported Luna operation: {name}")
    if name not in {"merge_cells", "unmerge_cells", "replace_grid"}:
        _rebuild_cells(candidate)
    return note


def _sync_component_grid(candidate: dict[str, Any], operation: dict[str, Any]) -> None:
    components = candidate.get("components") or []
    if not components or operation["op"] in {"set_metadata", "add_note"}:
        return
    if len(components) == 1:
        components[0]["grid"] = copy.deepcopy(candidate["grid"])
        return
    if operation["op"] != "set_cell":
        raise ValueError("Multi-component grid structure needs a reviewed reparse")
    rows = candidate["grid"]["rows"]
    row, column = _cell_position(operation["target"]["cell_id"])
    width = max((len(values) for values in rows), default=0)
    component_rows = [
        values + [""] * (width - len(values))
        for part in components for values in part["grid"]["rows"]
    ]
    original_rows = copy.deepcopy(rows)
    original_rows[row - 1][column - 1] = operation["before"]
    if component_rows != original_rows:
        raise ValueError("Component rows do not align with the logical grid")
    offset = 0
    for part in components:
        part_rows = part["grid"]["rows"]
        if row <= offset + len(part_rows):
            local_row = row - offset
            if column > len(part_rows[local_row - 1]):
                raise ValueError("Cell edit falls outside its component")
            if part_rows[local_row - 1][column - 1] != operation["before"]:
                raise ValueError("Component cell before mismatch")
            part_rows[local_row - 1][column - 1] = str(operation["after"])
            for cell in part["grid"].get("cells", []):
                if int(cell["row"]) == local_row and int(cell["column"]) == column:
                    cell["text"] = str(operation["after"])
            return
        offset += len(part_rows)
    raise ValueError("Cell edit falls outside all components")


def _validate_patch(
    candidate: dict[str, Any], patch: dict[str, Any], batch_id: str,
    table_dir: Path, contract_version: str,
) -> None:
    if contract_version != "2.0" or patch.get("schema_version") != 1 or patch.get("contract_version") != contract_version:
        raise ValueError("Patch schema/contract version mismatch")
    if patch.get("batch_id") != batch_id or patch.get("candidate_id") != candidate["candidate_id"]:
        raise ValueError("Patch batch_id or candidate_id mismatch")
    if Path(str(patch.get("table_directory", ""))).resolve() != table_dir.resolve():
        raise ValueError("Patch table_directory mismatch")
    decision = patch.get("decision")
    operations = patch.get("operations")
    unresolved = patch.get("unresolved")
    if decision not in {"accept", "patch", "partial"} or not isinstance(operations, list) or not isinstance(unresolved, list):
        raise ValueError("Invalid decision, operations, or unresolved")
    pages = sorted(int(value) for value in patch.get("reviewed_source_pages", []))
    if pages != sorted(int(value) for value in candidate.get("source_physical_pages", [])):
        raise ValueError("reviewed_source_pages must cover all candidate source pages")
    expected_codes = {
        issue["code"] for issue in candidate.get("issues", []) if issue.get("severity") in {"review", "error"}
    }
    resolutions = patch.get("issue_resolutions")
    if not isinstance(resolutions, list):
        raise ValueError("issue_resolutions must be a list")
    resolved_codes = {item.get("issue_code") for item in resolutions}
    if expected_codes != resolved_codes or len(resolutions) != len(resolved_codes):
        raise ValueError(f"issue_resolutions mismatch: expected {sorted(expected_codes)}, got {sorted(resolved_codes)}")
    if decision == "accept" and (operations or unresolved):
        raise ValueError("accept requires empty operations and unresolved")
    if decision == "patch" and (not operations or unresolved):
        raise ValueError("patch requires operations and no unresolved")
    if decision == "partial" and not unresolved:
        raise ValueError("partial requires unresolved entries")
    op_ids = [item.get("op_id") for item in operations]
    if len(op_ids) != len(set(op_ids)) or any(not value for value in op_ids):
        raise ValueError("Operation IDs must be present and unique")
    for operation in operations:
        if operation.get("confidence") not in {"high", "medium"}:
            raise ValueError("Every operation requires high or medium confidence")
        if not operation.get("evidence") or not operation.get("reason"):
            raise ValueError("Every operation requires evidence and reason")
    if contract_version == "2.0":
        source_pages = set(pages)
        for item in resolutions:
            if item.get("resolution") not in {"confirmed", "corrected", "false_positive", "unresolved"} or not str(item.get("reason") or "").strip():
                raise ValueError("Invalid issue resolution")
        if any(item["resolution"] == "unresolved" for item in resolutions) and decision != "partial":
            raise ValueError("Unresolved issue requires a partial decision")
        for operation in operations:
            for evidence in operation["evidence"]:
                if type(evidence.get("physical_page")) is not int or evidence["physical_page"] not in source_pages or not str(evidence.get("description") or "").strip():
                    raise ValueError("Invalid operation source-page evidence")
        for item in unresolved:
            if not item.get("code") or not item.get("description") or not item.get("consequence") or not isinstance(item.get("evidence_pages"), list) or any(type(page) is not int or page not in source_pages for page in item["evidence_pages"]):
                raise ValueError("Invalid unresolved entry")
        missing = {
            item["issue_code"] for item in resolutions if item["resolution"] == "unresolved"
        } - {item["code"] for item in unresolved}
        if missing:
            raise ValueError(f"Unresolved issues missing entries: {sorted(missing)}")


def _validate_formula_review(
    report: dict[str, Any], assigned: list[dict[str, Any]], batch_id: str
) -> tuple[set[int], list[dict[str, Any]]]:
    expected = {
        int(page)
        for item in assigned
        for page in item["source_physical_pages"]
    }
    reviews = report.get("formula_review")
    if not isinstance(reviews, list):
        raise ValueError(f"formula_review is missing in {batch_id}")
    if any(not isinstance(item, dict) or type(item.get("physical_page")) is not int for item in reviews):
        raise ValueError(f"Invalid formula_review page in {batch_id}")
    pages = [item["physical_page"] for item in reviews]
    if len(pages) != len(set(pages)) or set(pages) != expected:
        raise ValueError(f"formula_review page coverage mismatch in {batch_id}")
    exceptions: list[dict[str, Any]] = []
    candidate_pages = {
        item["candidate_id"]: set(item["source_physical_pages"]) for item in assigned
    }
    for item in reviews:
        status = item.get("status")
        findings = item.get("findings")
        if status not in {"no_formula", "clean", "issue", "uncertain"} or not isinstance(findings, list):
            raise ValueError(f"Invalid formula_review entry in {batch_id}")
        if (status == "issue") != bool(findings):
            raise ValueError(f"formula_review findings mismatch in {batch_id}")
        if status == "uncertain" and not item.get("note"):
            raise ValueError(f"formula_review uncertainty needs a note in {batch_id}")
        for finding in findings:
            if not isinstance(finding, dict) or any(
                not isinstance(finding.get(key), str) or not finding[key].strip()
                for key in ("location", "artifact", "source_visible", "suggested_latex", "kind")
            ) or not isinstance(finding.get("observed"), str) or "candidate_id" not in finding:
                raise ValueError(f"Invalid formula_review finding in {batch_id}")
            candidate_id = finding["candidate_id"]
            if candidate_id is not None and (
                candidate_id not in candidate_pages or item["physical_page"] not in candidate_pages[candidate_id]
            ):
                raise ValueError(f"Invalid formula_review finding in {batch_id}")
        if status in {"issue", "uncertain"}:
            exceptions.append({"batch_id": batch_id, **item})
    return expected, exceptions


def _refresh_identifier(candidate: dict[str, Any]) -> None:
    parsed = parse_identifier_prefix(str(candidate.get("node_id_raw") or ""))
    if parsed is None:
        candidate["node_id_core"] = None
        candidate["node_id_suffixes"] = []
        candidate["node_id_match_key"] = None
        return
    candidate["node_id_raw"] = parsed.raw
    candidate["node_id_core"] = parsed.core
    candidate["node_id_suffixes"] = list(parsed.suffixes)
    candidate["node_id_match_key"] = parsed.match_key


def _combine_candidates(run_dir: Path, candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    order: list[str] = []
    for candidate in candidates:
        _refresh_identifier(candidate)
        match_key = candidate.get("node_id_match_key")
        group_key = f"candidate:{candidate['candidate_id']}"
        if match_key:
            candidate_pages = set(candidate.get("source_physical_pages", []))
            candidate_title = re.sub(
                r"\s+", " ", str(candidate.get("caption") or "").casefold()
            ).strip()
            for existing_key in order:
                existing_parts = groups[existing_key]
                first = existing_parts[0]
                if first.get("node_id_match_key") != match_key:
                    continue
                existing_pages = {
                    page
                    for part in existing_parts
                    for page in part.get("source_physical_pages", [])
                }
                first_title = re.sub(
                    r"\s+", " ", str(first.get("caption") or "").casefold()
                ).strip()
                if candidate_pages & existing_pages and (
                    not candidate_title
                    or not first_title
                    or candidate_title == first_title
                ):
                    group_key = existing_key
                    break
            else:
                group_key = f"identified:{match_key}:{candidate['candidate_id']}"
        if group_key not in groups:
            groups[group_key] = []
            order.append(group_key)
        groups[group_key].append(candidate)
    logical: list[dict[str, Any]] = []
    for group_key in order:
        parts = groups[group_key]
        if len(parts) == 1 or any(p["data_status"] == "review_required" for p in parts):
            logical.extend(parts)
            continue
        identified = [part for part in parts if part.get("node_id_raw")]
        base = copy.deepcopy(identified[0] if identified else parts[0])
        combined_rows: list[list[str]] = []
        combined_cells: list[dict[str, Any]] = []
        combined_media: list[dict[str, Any]] = []
        combined_components: list[dict[str, Any]] = []
        row_offset = 0
        for part_index, part in enumerate(parts, 1):
            part_components = part.get("components") or [
                {
                    "component_id": "component_001",
                    "title_raw": "",
                    "fragment_ids": [
                        fragment.get("fragment_id")
                        for fragment in part.get("source_fragments", [])
                        if fragment.get("fragment_id")
                    ],
                    "source_physical_pages": part.get("source_physical_pages", []),
                    "grid": part["grid"],
                }
            ]
            for component in part_components:
                record = copy.deepcopy(component)
                record["component_id"] = f"component_{len(combined_components) + 1:03d}"
                combined_components.append(record)
            rows = copy.deepcopy(part["grid"]["rows"])
            combined_rows.extend(rows)
            for cell in part["grid"].get("cells", []):
                record = copy.deepcopy(cell)
                record["row"] = int(record["row"]) + row_offset
                record["cell_id"] = f"R{record['row']:03d}C{int(record['column']):03d}"
                record["source_subtable"] = part_index
                combined_cells.append(record)
            for media in part["grid"].get("embedded_media", []):
                record = copy.deepcopy(media)
                row, column = _cell_position(record["cell_id"])
                record["cell_id"] = f"R{row + row_offset:03d}C{column:03d}"
                combined_media.append(record)
            row_offset += len(rows)
        width = max((len(row) for row in combined_rows), default=0)
        for row in combined_rows:
            row.extend([""] * (width - len(row)))
        base["combined_candidates"] = [part["candidate_id"] for part in parts]
        base["source_physical_pages"] = sorted(
            {page for part in parts for page in part.get("source_physical_pages", [])}
        )
        base["source_fragments"] = [
            fragment for part in parts for fragment in part.get("source_fragments", [])
        ]
        base["footnotes"] = "\n".join(
            dict.fromkeys(part.get("footnotes", "") for part in parts if part.get("footnotes"))
        )
        base["grid"] = {
            "rows": combined_rows,
            "cells": combined_cells,
            "embedded_media": combined_media,
            "source_table_count": sum(
                int(part.get("grid", {}).get("source_table_count", 0)) for part in parts
            ),
            "format": "combined_html_tables",
        }
        base["components"] = combined_components
        base["issues"] = [issue for part in parts for issue in part.get("issues", [])]
        base["data_status"] = (
            "partial" if any(part.get("data_status") == "partial" for part in parts) else "luna_reviewed"
        )
        base["validation_route"] = "mineru_plus_rules_plus_luna_combined"
        destination = run_dir / base["artifact_directory"]
        source_artifacts = [part["artifact_directory"] for part in parts]
        base["combined_artifacts"] = [name for name in source_artifacts if name != base["artifact_directory"]]
        patches = [_read_json(run_dir / name / "luna_patch.json") for name in source_artifacts if (run_dir / name / "luna_patch.json").is_file()]
        if patches:
            _write_json(destination / "luna_patch.json", {"parts": patches})
        base["artifact_directory"] = destination.relative_to(run_dir).as_posix()
        _write_json(destination / "candidate.json", base)
        _write_workbook(destination / "table.xlsx", base)
        notes = _notes(base)
        notes += "\n## Combined MinerU blocks\n\n"
        notes += "This logical Table contains multiple source-visible submatrices that MinerU emitted as separate HTML table blocks. They are preserved sequentially in one workbook.\n"
        (destination / "notes.md").write_text(notes, encoding="utf-8")
        _render_candidate_preview(base, destination / "candidate_preview.png")
        _write_json(
            destination / "table_provenance.json",
            {
                "schema_version": 1,
                "manual_pdf": str(_read_json(run_dir / "manifest.json")["manual_pdf"]),
                "source_physical_pages": base["source_physical_pages"],
                "source_table_artifacts": source_artifacts,
                "validation_route": base["validation_route"],
                "luna_patch_applied": True,
                "unresolved_issues": [],
            },
        )
        logical.append(base)
    return logical


def apply_luna_results(*, run_dir: Path, force: bool = False) -> dict[str, Any]:
    run_dir = run_dir.expanduser().resolve()
    plan = _read_json(run_dir / "luna_batch_plan.json")
    results: list[dict[str, Any]] = []
    formula_pages: set[int] = set()
    formula_exceptions: list[dict[str, Any]] = []
    candidates = load_candidates(run_dir)
    revised_by_candidate: dict[str, dict[str, Any]] = {}
    pending = []
    for batch in plan["batches"]:
        if not (Path(batch["output_directory"]) / "batch_report.json").is_file():
            pending.append(batch["batch_id"])
            continue
        manifest = _read_json(Path(batch["input_manifest"]))
        output_dir = Path(batch["output_directory"])
        report = _read_json(output_dir / "batch_report.json")
        assigned = [item["candidate_id"] for item in manifest["assigned_candidates"]]
        contract_version = manifest.get("contract_version")
        if contract_version != "2.0" or report.get("contract_version") != "2.0":
            raise ValueError("Only the workbench Luna contract 2.0 is supported")
        if report.get("schema_version") != 1 or report.get("batch_id") != batch["batch_id"]:
            raise ValueError(f"Batch report mismatch: {batch['batch_id']}")
        pages, exceptions = _validate_formula_review(report, manifest["assigned_candidates"], batch["batch_id"])
        formula_pages.update(pages)
        formula_exceptions.extend(exceptions)
        if batch.get("applied"):
            continue
        report_items = report.get("results", [])
        report_map = {item["candidate_id"]: item for item in report_items}
        if len(report_map) != len(report_items) or set(report_map) != set(assigned):
            raise ValueError(f"Batch report candidate mismatch: {batch['batch_id']}")
        for item in manifest["assigned_candidates"]:
            table_dir = Path(item["table_directory"])
            candidate = _read_json(table_dir / "candidate.json")
            patch_path = (output_dir / report_map[item["candidate_id"]]["patch_file"]).resolve()
            if not patch_path.is_relative_to(output_dir.resolve()):
                raise ValueError("Patch path escapes its batch")
            patch = _read_json(patch_path)
            applied_path = table_dir / "luna_patch.json"
            if applied_path.is_file():
                if _read_json(applied_path) != patch:
                    raise ValueError("Patch changed after application; reconcile explicitly")
                revised_by_candidate[candidate["candidate_id"]] = candidate
                continue
            if item["base_revision"] != candidate["revision"]:
                raise ValueError("Candidate changed after Luna was assigned")
            _validate_patch(candidate, patch, batch["batch_id"], table_dir, contract_version)
            revised = copy.deepcopy(candidate)
            added_notes: list[str] = []
            for operation in patch["operations"]:
                if note := _apply_operation(revised, operation):
                    added_notes.append(note)
                _sync_component_grid(revised, operation)
            revised["data_status"] = "partial" if patch["decision"] == "partial" else "luna_reviewed"
            revised["validation_route"] = "mineru_plus_rules_plus_luna"
            _refresh_identifier(revised)
            destination = table_dir
            revised["artifact_directory"] = destination.relative_to(run_dir).as_posix()
            _write_workbook(destination / "table.xlsx", revised)
            notes = _notes(revised)
            if added_notes:
                notes += "\n## Luna review notes\n\n" + "\n\n".join(added_notes) + "\n"
            if patch.get("review_notes"):
                notes += "\n## Review summary\n\n" + str(patch["review_notes"]) + "\n"
            (destination / "notes.md").write_text(notes, encoding="utf-8")
            _render_candidate_preview(revised, destination / "candidate_preview.png")
            provenance_path = destination / "table_provenance.json"
            provenance = _read_json(provenance_path)
            provenance["validation_route"] = revised["validation_route"]
            provenance["luna_patch_applied"] = bool(patch["operations"])
            provenance["luna_decision"] = patch["decision"]
            provenance["unresolved_issues"] = patch["unresolved"]
            provenance["luna_patch_file"] = "luna_patch.json"
            _write_json(provenance_path, provenance)
            _write_json(destination / "candidate.json", revised)
            shutil.copy2(patch_path, destination / "luna_patch.json")
            revised_by_candidate[candidate["candidate_id"]] = revised
            results.append(
                {
                    "candidate_id": candidate["candidate_id"],
                    "decision": patch["decision"],
                    "artifact_directory": revised["artifact_directory"],
                }
            )

        batch["applied"] = True
        _write_json(run_dir / "luna_batch_plan.json", plan)
    candidates = [
        revised_by_candidate.get(candidate["candidate_id"], candidate)
        for candidate in candidates
    ]
    candidates = _combine_candidates(run_dir, candidates)
    write_index(run_dir, candidates)
    for candidate in candidates:
        for name in candidate.get("combined_artifacts", []):
            path = (run_dir / name).resolve()
            if path.parent == run_dir and path.name.startswith("Table_") and path.exists():
                shutil.rmtree(path)
    manifest = _read_json(run_dir / "manifest.json")
    by_key: dict[str, list[dict[str, Any]]] = {}
    by_synthetic: dict[str, dict[str, Any]] = {}
    for candidate in candidates:
        by_key.setdefault(candidate.get("node_id_match_key") or "", []).append(candidate)
        if candidate.get("synthetic_id"):
            by_synthetic[candidate["synthetic_id"]] = candidate
    for node in manifest["nodes"]:
        if node.get("node_type") != "table":
            continue
        candidate = by_synthetic.get(node.get("synthetic_id") or "")
        if candidate is None:
            options = by_key.get(node.get("node_id_match_key") or "", [])
            exact = [item for item in options if (item.get("node_id_raw") or "").casefold() == (node.get("node_id_raw") or "").casefold()]
            candidate = exact[0] if len(exact) == 1 else options[0] if len(options) == 1 else None
        if candidate:
            node["data_status"] = candidate["data_status"]
            node["table_artifact"] = candidate["artifact_directory"]
            node["source_physical_pages"] = candidate["source_physical_pages"]
    _write_json(run_dir / "manifest.json", manifest)
    counts = Counter(item["decision"] for item in results)
    summary = {
        "reviewed_tables": sum(c["data_status"] != "review_required" for c in candidates),
        "pending_batches": pending,
        "decisions": dict(counts),
        "formula_pages_reviewed": len(formula_pages),
        "formula_exceptions": formula_exceptions,
        "final_manifest": str((run_dir / "manifest.json").resolve()),
    }
    _write_json(run_dir / "luna_apply_summary.json", summary)
    return summary
