from __future__ import annotations

from materials_workbench.storage import write_json as _atomic_json, publish_candidate, write_index

from collections import Counter, defaultdict
import json
from pathlib import Path
import re
from typing import Any, Iterable

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from .identifiers import ParsedCaption, parse_caption
from .page_sources import indexed_page_sources, validate_single_page_sources


def _write_json(path: Path, value: Any) -> None:
    if path.name == "candidate.json":
        publish_candidate(path.parent, value)
    else:
        _atomic_json(path, value)


def _textish(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return " ".join(filter(None, (_textish(item) for item in value))).strip()
    if isinstance(value, dict):
        for key in ("content", "text", "caption", "value"):
            if key in value:
                return _textish(value[key])
    return str(value).strip()


def _walk_blocks(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, list):
        for item in value:
            yield from _walk_blocks(item)
    elif isinstance(value, dict):
        if "type" in value:
            yield value
        else:
            for child in value.values():
                if isinstance(child, (list, dict)):
                    yield from _walk_blocks(child)


def _table_html(block: dict[str, Any]) -> str:
    for key in ("table_body", "html", "table_html", "content"):
        value = block.get(key)
        if isinstance(value, str) and (
            "<table" in value.casefold() or "|" in value
        ):
            return value
    return ""


def _caption_from_block(block: dict[str, Any]) -> tuple[str, ParsedCaption | None]:
    pieces = []
    for key in ("table_caption", "caption", "title"):
        value = _textish(block.get(key))
        if value:
            pieces.append(value)
    caption = " ".join(pieces).strip()
    parsed = parse_caption(caption, "table") if caption else None
    return caption, parsed


def _caption_in_text(block: dict[str, Any]) -> tuple[str, ParsedCaption | None]:
    text = _textish(block.get("text") or block.get("content"))
    for line in text.splitlines():
        # Arbitrary-position searches turn "See Table 3.2..." into a pending
        # caption for the next OCR table block. Only line-leading captions are
        # eligible here.
        if not re.match(r"(?i)^\s*(?:Table|Tbl\.?)(?![A-Za-z])", line):
            continue
        candidate = line.strip()
        parsed = parse_caption(candidate, "table")
        if parsed:
            return candidate, parsed
    return "", None


def _page_index(block: dict[str, Any]) -> int | None:
    for key in ("page_idx", "page_index", "page"):
        value = block.get(key)
        if isinstance(value, int):
            return value
        if isinstance(value, str) and value.isdigit():
            return int(value)
    return None


def _bbox(block: dict[str, Any]) -> list[float] | None:
    value = block.get("bbox") or block.get("table_bbox")
    if isinstance(value, list) and len(value) == 4:
        try:
            return [float(item) for item in value]
        except (TypeError, ValueError):
            return None
    return None


def _looks_like_standalone_table_title(text: str) -> bool:
    value = re.sub(r"\s+", " ", text).strip()
    if not value or len(value) > 120 or value.endswith((".", ";", ":")):
        return False
    if re.match(r"(?i)^(?:notes?|see|source|condition|alloy|form|footnotes?)\b", value):
        return False
    if re.match(r"(?i)^(?:AMS|ASTM|SAE|MIL|ISO|DIN|EN)\s*[-:]?\s*[A-Z0-9./-]+(?:\s+.*)?$", value):
        return True
    letters = [character for character in value if character.isalpha()]
    return (
        bool(letters)
        and any(character.isdigit() for character in value)
        and len(value.split()) <= 10
        and sum(character.isupper() for character in letters) / len(letters) >= 0.85
    )


def _is_tightly_above(title_block: dict[str, Any], table_block: dict[str, Any]) -> bool:
    title_bbox = title_block.get("bbox")
    table_bbox = _bbox(table_block)
    if not title_bbox or not table_bbox:
        return True
    gap = table_bbox[1] - title_bbox[3]
    # MinerU bboxes may overlap a title by a few pixels even when the title is
    # visibly above the table border.
    return -20 <= gap <= 100


def _normalized_header(grid: dict[str, Any]) -> str:
    rows = grid.get("rows") or []
    if not rows:
        return ""
    return re.sub(
        r"[^0-9a-z]+",
        " ",
        " | ".join(str(value) for value in rows[0]).casefold(),
    ).strip()


def _looks_like_cross_page_continuation(
    previous: dict[str, Any], current_html: str, current_bbox: list[float] | None
) -> bool:
    previous_grid = previous.get("grid") or {}
    current_grid = normalize_table_html(current_html)
    previous_width = max((len(row) for row in previous_grid.get("rows", [])), default=0)
    current_width = max((len(row) for row in current_grid.get("rows", [])), default=0)
    if not previous_width or previous_width != current_width:
        return False
    previous_header = _normalized_header(previous_grid)
    current_header = _normalized_header(current_grid)
    if not previous_header or previous_header != current_header:
        return False
    previous_bbox = previous.get("bbox")
    return bool(
        previous_bbox
        and current_bbox
        and previous_bbox[3] >= 700
        and current_bbox[1] <= 250
    )


def _physical_pages_from_block(
    block: dict[str, Any], page_map: dict[str, int]
) -> list[int]:
    local_pages: set[int] = set()
    start = _page_index(block)
    if start is not None:
        local_pages.add(start)
    for key in ("pages", "page_indices", "page_idxs", "page_range"):
        value = block.get(key)
        if isinstance(value, list):
            local_pages.update(int(item) for item in value if str(item).isdigit())
    start_value = block.get("start_page_idx")
    end_value = block.get("end_page_idx")
    if str(start_value).isdigit() and str(end_value).isdigit():
        local_pages.update(range(int(start_value), int(end_value) + 1))
    return sorted(
        {
            int(page_map[str(local_page)])
            for local_page in local_pages
            if str(local_page) in page_map
        }
    )


def _markdown_grid(text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in text.splitlines():
        if "|" not in line:
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def normalize_table_html(html: str) -> dict[str, Any]:
    soup = BeautifulSoup(html, "lxml")
    tables = soup.find_all("table")
    if not tables:
        rows = _markdown_grid(html)
        return {
            "rows": rows,
            "cells": [
                {
                    "cell_id": f"R{row_index:03d}C{column_index:03d}",
                    "row": row_index,
                    "column": column_index,
                    "rowspan": 1,
                    "colspan": 1,
                    "text": value,
                }
                for row_index, row in enumerate(rows, 1)
                for column_index, value in enumerate(row, 1)
            ],
            "embedded_media": [],
            "source_table_count": 0,
            "format": "markdown",
        }

    combined_rows: list[list[str]] = []
    combined_cells: list[dict[str, Any]] = []
    embedded: list[dict[str, Any]] = []
    row_offset = 0
    for table_number, table in enumerate(tables, 1):
        occupied: dict[tuple[int, int], bool] = {}
        local_rows: list[list[str]] = []
        local_cells: list[dict[str, Any]] = []
        for local_row, tr in enumerate(table.find_all("tr", recursive=True), 1):
            while len(local_rows) < local_row:
                local_rows.append([])
            column = 1
            direct_cells = tr.find_all(["th", "td"], recursive=False)
            if not direct_cells:
                direct_cells = tr.find_all(["th", "td"])
            for cell in direct_cells:
                while occupied.get((local_row, column)):
                    column += 1
                try:
                    rowspan = max(1, int(cell.get("rowspan", 1)))
                except (TypeError, ValueError):
                    rowspan = 1
                try:
                    colspan = max(1, int(cell.get("colspan", 1)))
                except (TypeError, ValueError):
                    colspan = 1
                text = " ".join(cell.stripped_strings)
                global_row = row_offset + local_row
                record = {
                    "cell_id": f"R{global_row:03d}C{column:03d}",
                    "row": global_row,
                    "column": column,
                    "rowspan": rowspan,
                    "colspan": colspan,
                    "text": text,
                    "source_table": table_number,
                }
                local_cells.append(record)
                for delta_row in range(rowspan):
                    target_local_row = local_row + delta_row
                    while len(local_rows) < target_local_row:
                        local_rows.append([])
                    for delta_column in range(colspan):
                        target_column = column + delta_column
                        occupied[(target_local_row, target_column)] = True
                        while len(local_rows[target_local_row - 1]) < target_column:
                            local_rows[target_local_row - 1].append("")
                        if delta_row == 0 and delta_column == 0:
                            local_rows[target_local_row - 1][target_column - 1] = text
                for tag in cell.find_all(["img", "svg", "math"]):
                    embedded.append(
                        {
                            "cell_id": record["cell_id"],
                            "tag": tag.name,
                            "source": tag.get("src") if tag.name == "img" else str(tag)[:500],
                        }
                    )
                column += colspan
        width = max((len(row) for row in local_rows), default=0)
        for row in local_rows:
            row.extend([""] * (width - len(row)))
        # MinerU may emit one HTML table per source page.  Remove only a fully
        # identical repeated header row; all other continuation rows remain.
        if combined_rows and local_rows and local_rows[0] == combined_rows[0]:
            local_rows = local_rows[1:]
            for cell in local_cells:
                cell["row"] -= 1
                cell["cell_id"] = f"R{cell['row']:03d}C{cell['column']:03d}"
            local_cells = [cell for cell in local_cells if cell["row"] > row_offset]
        combined_rows.extend(local_rows)
        combined_cells.extend(local_cells)
        row_offset = len(combined_rows)
    width = max((len(row) for row in combined_rows), default=0)
    for row in combined_rows:
        row.extend([""] * (width - len(row)))
    return {
        "rows": combined_rows,
        "cells": combined_cells,
        "embedded_media": embedded,
        "source_table_count": len(tables),
        "format": "html",
    }


def _issue(code: str, message: str, severity: str, **extra: Any) -> dict[str, Any]:
    return {"code": code, "message": message, "severity": severity, **extra}


def quality_issues(candidate: dict[str, Any]) -> list[dict[str, Any]]:
    grid = candidate["grid"]
    rows = grid["rows"]
    issues: list[dict[str, Any]] = []
    if candidate["node_id_raw"] is None:
        if candidate.get("identifier_status") == "absent_in_source":
            issues.append(
                _issue(
                    "unnumbered_table",
                    "The source shows a standalone Table without a formal Table number.",
                    "review",
                    synthetic_id=candidate.get("synthetic_id"),
                )
            )
        else:
            issues.append(_issue("unresolved_identifier", "No Table identifier was recovered.", "review"))
    if not candidate["source_physical_pages"]:
        issues.append(_issue("missing_page_provenance", "No original physical page is associated.", "error"))
    if not rows or not any(any(cell.strip() for cell in row) for row in rows):
        issues.append(_issue("empty_grid", "No non-empty table cells were recovered.", "error"))
    widths = {len(row) for row in rows}
    if len(widths) > 1:
        issues.append(_issue("inconsistent_width", "Rows have inconsistent widths.", "review"))
    fragment_widths = {
        width for width in candidate.get("fragment_column_counts", []) if width > 0
    }
    if len(fragment_widths) > 1:
        issues.append(
            _issue(
                "cross_fragment_column_mismatch",
                "MinerU table fragments disagree on the logical column count.",
                "review",
                column_counts=sorted(fragment_widths),
            )
        )
    if grid["source_table_count"] > max(1, len(candidate["source_physical_pages"])):
        issues.append(
            _issue(
                "multiple_html_tables",
                "More HTML table fragments than source pages were returned.",
                "review",
            )
        )
    if grid["embedded_media"]:
        issues.append(
            _issue(
                "embedded_media",
                "One or more cells contain an image, SVG, or formula element.",
                "review",
                cells=sorted({item["cell_id"] for item in grid["embedded_media"]}),
            )
        )
    association_reasons = {
        fragment.get("association_reason")
        for fragment in candidate.get("source_fragments", [])
    }
    if "same_page_component_without_caption" in association_reasons:
        issues.append(
            _issue(
                "same_page_multiple_components",
                "A second same-page table block was associated as another component of this logical Table.",
                "review",
            )
        )
    if "orphan_table_fragment_preserved" in association_reasons:
        issues.append(
            _issue(
                "orphan_table_fragment",
                "An uncaptioned table block was preserved independently because continuation evidence was insufficient.",
                "review",
            )
        )
    large_merges = [
        cell["cell_id"]
        for cell in grid["cells"]
        if cell["rowspan"] > 3 or cell["colspan"] > 5
    ]
    if large_merges:
        issues.append(
            _issue(
                "large_merged_region",
                "Unusually large merged regions need visual confirmation.",
                "review",
                cells=large_merges,
            )
        )
    suspicious: list[str] = []
    for cell in grid["cells"]:
        text = cell["text"]
        if re.search(r"\d[OoIl]\d|\d[OoIl](?:\.|,)|(?:\.|,)[OoIl]\d", text):
            suspicious.append(cell["cell_id"])
    if suspicious:
        issues.append(
            _issue(
                "suspicious_numeric_ocr",
                "Letter-like glyphs occur inside numeric-looking values.",
                "review",
                cells=suspicious,
            )
        )
    if len(candidate["source_physical_pages"]) > 1:
        issues.append(
            _issue(
                "multi_page_table",
                "The logical table spans multiple source pages; provenance was preserved.",
                "info",
            )
        )
    return issues


def _safe_key(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-z._()\-]+", "_", value).strip("._")
    return value or "unresolved"


def _write_workbook(path: Path, candidate: dict[str, Any]) -> None:
    workbook = Workbook()
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def populate(sheet, grid: dict[str, Any]) -> None:
        rows = grid["rows"]
        for row_index, row in enumerate(rows, 1):
            for column_index, value in enumerate(row, 1):
                cell = sheet.cell(row_index, column_index, _excel_value(value))
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
                if row_index == 1:
                    cell.font = Font(name="Calibri", size=11, bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in grid["cells"]:
            if cell["rowspan"] > 1 or cell["colspan"] > 1:
                start_row = cell["row"]
                start_column = cell["column"]
                end_row = start_row + cell["rowspan"] - 1
                end_column = start_column + cell["colspan"] - 1
                if end_row <= max(1, len(rows)):
                    sheet.merge_cells(
                        start_row=start_row,
                        start_column=start_column,
                        end_row=end_row,
                        end_column=end_column,
                    )
        for column_index, column_cells in enumerate(sheet.columns, 1):
            letter = get_column_letter(column_index)
            sheet.column_dimensions[letter].width = min(
                45,
                max(
                    10,
                    max(
                        (len(str(cell.value or "")) for cell in column_cells),
                        default=0,
                    )
                    + 2,
                ),
            )

    components = candidate.get("components") or []
    if not components:
        components = [{"component_id": "component_001", "grid": candidate["grid"]}]
    for index, component in enumerate(components, 1):
        sheet = workbook.active if index == 1 else workbook.create_sheet()
        sheet.title = "Table" if len(components) == 1 else f"Component {index}"
        populate(sheet, component["grid"])
    metadata = workbook.create_sheet("Provenance")
    metadata.append(["field", "value"])
    metadata.append(["node_id_raw", candidate["node_id_raw"] or ""])
    metadata.append(["caption", candidate["caption"]])
    metadata.append(
        ["source_physical_pages", ", ".join(map(str, candidate["source_physical_pages"]))]
    )
    metadata.append(["data_status", candidate["data_status"]])
    metadata.append(["component_count", len(components)])
    for component in components:
        metadata.append(
            [
                component["component_id"],
                "pages="
                + ",".join(map(str, component.get("source_physical_pages", [])))
                + "; fragments="
                + ",".join(component.get("fragment_ids", [])),
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _excel_value(value: str) -> str | int | float:
    stripped = value.strip()
    if re.fullmatch(r"[+-]?\d+", stripped):
        unsigned = stripped.lstrip("+-")
        if len(unsigned) == 1 or not unsigned.startswith("0"):
            return int(stripped)
    if re.fullmatch(r"[+-]?(?:\d+\.\d*|\d*\.\d+)(?:[Ee][+-]?\d+)?", stripped):
        return float(stripped)
    return value


def _render_candidate_preview(candidate: dict[str, Any], target: Path) -> None:
    rows = candidate["grid"]["rows"]
    width = max((len(row) for row in rows), default=1)
    cell_width, cell_height = 180, 42
    image = Image.new(
        "RGB",
        (max(1, width) * cell_width + 1, max(1, len(rows)) * cell_height + 1),
        "white",
    )
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            left, top = column_index * cell_width, row_index * cell_height
            draw.rectangle(
                (left, top, left + cell_width, top + cell_height), outline="#808080"
            )
            value = row[column_index] if column_index < len(row) else ""
            clipped = value.replace("\n", " ")[:48]
            draw.text((left + 4, top + 4), clipped, fill="black", font=font)
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)






def _notes(candidate: dict[str, Any]) -> str:
    lines = [f"# Table {candidate['node_id_raw'] or 'unresolved'}", ""]
    if candidate["caption"]:
        lines.extend([candidate["caption"], ""])
    if candidate["footnotes"]:
        lines.extend(["## Footnotes", "", candidate["footnotes"], ""])
    if candidate["grid"]["embedded_media"]:
        lines.extend(
            [
                "## Omitted cell media",
                "",
                "One or more cells contain image/formula markup. The two-dimensional "
                "candidate preserves the surrounding text, while the media requires review.",
                "",
            ]
        )
    lines.extend(
        [
            "## Provenance",
            "",
            "Original physical pages: "
            + ", ".join(map(str, candidate["source_physical_pages"])),
            "",
            f"Data status: `{candidate['data_status']}`",
            "",
        ]
    )
    return "\n".join(lines)


def _review_request(candidate: dict[str, Any]) -> dict[str, Any]:
    grid = candidate["grid"]
    return {
        "schema_version": 1,
        "candidate_id": candidate["candidate_id"],
        "node_id_raw": candidate["node_id_raw"],
        "synthetic_id": candidate.get("synthetic_id"),
        "caption": candidate["caption"],
        "footnotes": candidate["footnotes"],
        "source_physical_pages": candidate["source_physical_pages"],
        "source_fragments": candidate.get("source_fragments", []),
        "rows": grid["rows"],
        "merged_cells": [
            {key: cell[key] for key in ("cell_id", "rowspan", "colspan")}
            for cell in grid["cells"]
            if cell["rowspan"] > 1 or cell["colspan"] > 1
        ],
        "embedded_media": grid.get("embedded_media", []),
        "components": [
            {
                "component_id": part["component_id"],
                "title_raw": part["title_raw"],
                "fragment_ids": part["fragment_ids"],
                "source_physical_pages": part["source_physical_pages"],
                "row_count": len(part["grid"]["rows"]),
                "column_count": max((len(row) for row in part["grid"]["rows"]), default=0),
            }
            for part in candidate.get("components", [])
        ],
        "issues": candidate.get("issues", []),
    }


def _merge_manifest_tables(
    manual_manifest: dict[str, Any], candidates: list[dict[str, Any]]
) -> dict[str, Any]:
    nodes = manual_manifest["nodes"]
    by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in nodes:
        node_key = node.get("node_id_match_key") or node.get("identifier_match_key")
        if node["node_type"] == "table" and node_key:
            by_key[node_key].append(node)
    sections = [node for node in nodes if node["node_type"] == "section"]
    next_sequence = max((node["sequence"] for node in nodes), default=0) + 1
    for candidate in candidates:
        key = candidate.get("node_id_match_key")
        matches = by_key.get(key, []) if key else []
        exact = [
            node
            for node in matches
            if node["node_id_raw"].casefold() == (candidate["node_id_raw"] or "").casefold()
        ]
        if len(exact) == 1 or (len(matches) == 1 and not exact):
            node = exact[0] if exact else matches[0]
            node["node_id_core"] = candidate["node_id_core"]
            node["node_id_suffixes"] = candidate["node_id_suffixes"]
            node["node_id_match_key"] = candidate["node_id_match_key"]
            node["source_physical_pages"] = candidate["source_physical_pages"]
            node["data_status"] = candidate["data_status"]
            node["table_artifact"] = candidate["artifact_directory"]
            continue
        if not candidate["node_id_raw"] and not candidate.get("synthetic_id"):
            continue
        core = candidate["node_id_core"]
        parent = None
        if core:
            prefixes = [".".join(core.split(".")[:index]) for index in range(len(core.split(".")), 0, -1)]
            for prefix in prefixes:
                parent = next(
                    (
                        item
                        for item in reversed(sections)
                        if (item.get("node_id_core") or item.get("identifier_core"))
                        == prefix
                    ),
                    None,
                )
                if parent:
                    break
        elif candidate["source_physical_pages"]:
            first_page = min(candidate["source_physical_pages"])
            preceding_leaves = [
                item
                for item in nodes
                if item["node_type"] in {"table", "figure"}
                and item.get("parent_sequence") is not None
                and (
                    item.get("target_page")
                    or min(item.get("source_physical_pages") or [10**9])
                )
                <= first_page
            ]
            if preceding_leaves:
                anchor = max(
                    preceding_leaves,
                    key=lambda item: (
                        item.get("target_page")
                        or max(item.get("source_physical_pages") or [0]),
                        item["sequence"],
                    ),
                )
                parent = next(
                    (
                        item
                        for item in sections
                        if item["sequence"] == anchor["parent_sequence"]
                    ),
                    None,
                )
        node = {
            "sequence": next_sequence,
            "node_type": "table",
            "node_id_raw": candidate["node_id_raw"],
            "synthetic_id": candidate.get("synthetic_id"),
            "identifier_status": candidate.get("identifier_status", "present"),
            "node_id_core": candidate["node_id_core"],
            "node_id_suffixes": candidate["node_id_suffixes"],
            "node_id_match_key": key,
            "title_raw": candidate["caption"],
            "parent_sequence": parent["sequence"] if parent else None,
            "catalog_pages": [],
            "source_physical_pages": candidate["source_physical_pages"],
            "target_page": min(candidate["source_physical_pages"], default=None),
            "data_status": candidate["data_status"],
            "table_artifact": candidate["artifact_directory"],
            "discovered_by": "mineru",
        }
        nodes.append(node)
        if key:
            by_key[key].append(node)
        next_sequence += 1
    return manual_manifest


def process_mineru_results(*, run_dir: Path, force: bool = False) -> dict[str, Any]:
    run_dir = run_dir.expanduser().resolve()
    segments = json.loads((run_dir / "segments.json").read_text(encoding="utf-8"))
    manual_manifest = json.loads(
        (run_dir / "manifest.json").read_text(encoding="utf-8")
    )
    catalog_by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in manual_manifest["nodes"]:
        node_key = node.get("node_id_match_key") or node.get("identifier_match_key")
        if node["node_type"] == "table" and node_key:
            catalog_by_key[node_key].append(node)
    pdf_path = Path(manual_manifest["manual_pdf"])
    page_review_dir = Path(manual_manifest["page_review_dir"])
    page_sources = indexed_page_sources(
        page_review_dir, expected_page_count=manual_manifest["page_count"]
    )
    validate_single_page_sources(page_sources)
    fragments: list[dict[str, Any]] = []
    unresolved_counter = 0
    for segment in segments:
        content_path = run_dir / "ocr" / segment["segment_id"] / "content_list.json"
        if not content_path.is_file():
            continue
        content = json.loads(content_path.read_text(encoding="utf-8"))
        active_key: str | None = None
        active_last_page: int | None = None
        active_last_fragment: dict[str, Any] | None = None
        pending_caption: dict[str, Any] | None = None
        last_text_by_page: dict[int, dict[str, Any]] = {}
        unnumbered_by_page: Counter[int] = Counter()
        for block_index, block in enumerate(_walk_blocks(content)):
            block_type = str(block.get("type", "")).casefold()
            html = _table_html(block)
            page_map = segment["segment_page_to_physical_page"]
            physical_pages = _physical_pages_from_block(block, page_map)
            if "table" not in block_type and not html:
                text_caption, text_parsed = _caption_in_text(block)
                if text_parsed:
                    pending_caption = {
                        "caption": text_caption,
                        "parsed": text_parsed,
                        "physical_pages": physical_pages,
                        "block_index": block_index,
                        "bbox": _bbox(block),
                    }
                text_value = _textish(block.get("text") or block.get("content"))
                if physical_pages and text_value and block_type in {"text", "title", "heading"}:
                    last_text_by_page[physical_pages[0]] = {
                        "text": text_value,
                        "block_index": block_index,
                        "bbox": _bbox(block),
                    }
                continue
            caption, parsed = _caption_from_block(block)
            if (
                parsed is None
                and pending_caption is not None
                and physical_pages
                and pending_caption["physical_pages"]
                and pending_caption["physical_pages"][0] == physical_pages[0]
                and pending_caption["block_index"] < block_index
            ):
                caption = pending_caption["caption"]
                parsed = pending_caption["parsed"]
            pending_caption = None
            first_page = physical_pages[0] if physical_pages else None
            preceding_text = last_text_by_page.get(first_page) if first_page else None
            standalone_title = ""
            if (
                parsed is None
                and preceding_text is not None
                and preceding_text["block_index"] < block_index
                and block_index - preceding_text["block_index"] <= 2
                and _looks_like_standalone_table_title(preceding_text["text"])
                and _is_tightly_above(preceding_text, block)
            ):
                standalone_title = preceding_text["text"]
            current_grid = normalize_table_html(html)
            association_reason = ""
            synthetic_id: str | None = None
            if parsed:
                active_key = (
                    segment["segment_id"]
                    + "::"
                    + parsed.identifier.match_key
                    + "::"
                    + parsed.identifier.raw.casefold()
                )
                association_reason = "explicit_table_caption"
            elif standalone_title and first_page is not None:
                unnumbered_by_page[first_page] += 1
                synthetic_id = (
                    f"unnumbered-p{first_page:04d}-{unnumbered_by_page[first_page]:02d}"
                )
                active_key = f"{segment['segment_id']}::{synthetic_id}"
                association_reason = "standalone_title_without_table_number"
            elif active_key is not None and first_page == active_last_page:
                association_reason = "same_page_component_without_caption"
            elif (
                active_key is not None
                and first_page is not None
                and active_last_page is not None
                and first_page == active_last_page + 1
                and active_last_fragment is not None
                and _looks_like_cross_page_continuation(
                    active_last_fragment, html, _bbox(block)
                )
            ):
                association_reason = "compatible_repeated_header_continuation"
            else:
                unresolved_counter += 1
                if first_page is not None:
                    unnumbered_by_page[first_page] += 1
                    synthetic_id = (
                        f"unnumbered-p{first_page:04d}-{unnumbered_by_page[first_page]:02d}"
                    )
                    active_key = f"{segment['segment_id']}::{synthetic_id}"
                else:
                    active_key = f"unresolved:{segment['segment_id']}:{unresolved_counter:03d}"
                association_reason = "orphan_table_fragment_preserved"
            if physical_pages:
                active_last_page = max(physical_pages)
            fragment = {
                "fragment_id": f"{segment['segment_id']}::block_{block_index:04d}",
                "group_key": active_key,
                "segment_id": segment["segment_id"],
                "block_index": block_index,
                "physical_pages": physical_pages,
                "caption": caption,
                "parsed": parsed,
                "synthetic_id": synthetic_id,
                "identifier_status": (
                    "present" if parsed else "absent_in_source" if synthetic_id else "unresolved"
                ),
                "standalone_title": standalone_title,
                "association_reason": association_reason,
                "footnotes": _textish(block.get("table_footnote") or block.get("footnote")),
                "html": html,
                "bbox": _bbox(block),
                "image_path": block.get("img_path") or block.get("image_path"),
                "grid": current_grid,
            }
            fragments.append(fragment)
            active_last_fragment = fragment

    _write_json(
        run_dir / "table_fragments.json",
        [
            {
                key: value
                for key, value in fragment.items()
                if key not in {"parsed", "grid"}
            }
            | {
                "node_id_raw": (
                    fragment["parsed"].identifier.raw if fragment["parsed"] else None
                ),
                "node_id_match_key": (
                    fragment["parsed"].identifier.match_key if fragment["parsed"] else None
                ),
                "grid_shape": {
                    "rows": len(fragment["grid"]["rows"]),
                    "columns": max(
                        (len(row) for row in fragment["grid"]["rows"]), default=0
                    ),
                },
            }
            for fragment in fragments
        ],
    )

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for fragment in fragments:
        grouped[fragment["group_key"]].append(fragment)
    identifier_group_counts: Counter[str] = Counter()
    for parts in grouped.values():
        parsed = next((part["parsed"] for part in parts if part["parsed"]), None)
        if parsed:
            identifier_group_counts[parsed.identifier.match_key] += 1
    candidates: list[dict[str, Any]] = []
    for group_key, parts in grouped.items():
        parsed = next((part["parsed"] for part in parts if part["parsed"]), None)
        synthetic_id = next(
            (part["synthetic_id"] for part in parts if part.get("synthetic_id")), None
        )
        standalone_title = next(
            (
                part["standalone_title"]
                for part in parts
                if part.get("standalone_title")
            ),
            "",
        )
        caption = next((part["caption"] for part in parts if part["caption"]), "")
        if not caption:
            caption = standalone_title
        source_pages = sorted(
            {
                int(physical_page)
                for part in parts
                for physical_page in part["physical_pages"]
            }
        )
        combined_html = "\n".join(part["html"] for part in parts if part["html"])
        grid = normalize_table_html(combined_html)
        fragment_column_counts = []
        for part in parts:
            part_grid = normalize_table_html(part["html"])
            fragment_column_counts.append(
                max((len(row) for row in part_grid["rows"]), default=0)
            )
        node_id = parsed.identifier.raw if parsed else None
        components: list[dict[str, Any]] = []
        for part in parts:
            start_new_component = (
                not components
                or part.get("association_reason")
                in {
                    "same_page_component_without_caption",
                    "standalone_title_without_table_number",
                    "orphan_table_fragment_preserved",
                }
            )
            if start_new_component:
                components.append(
                    {
                        "component_id": f"component_{len(components) + 1:03d}",
                        "title_raw": part.get("standalone_title") or "",
                        "fragment_ids": [],
                        "source_physical_pages": [],
                        "html_parts": [],
                    }
                )
            component = components[-1]
            component["fragment_ids"].append(part["fragment_id"])
            component["source_physical_pages"] = sorted(
                set(component["source_physical_pages"]) | set(part["physical_pages"])
            )
            component["html_parts"].append(part["html"])
        for component in components:
            component["grid"] = normalize_table_html("\n".join(component.pop("html_parts")))
        candidate = {
            "candidate_id": group_key,
            "node_id_raw": node_id,
            "node_id_core": parsed.identifier.core if parsed else None,
            "node_id_suffixes": list(parsed.identifier.suffixes) if parsed else [],
            "node_id_match_key": parsed.identifier.match_key if parsed else None,
            "synthetic_id": synthetic_id,
            "identifier_status": (
                "present" if parsed else "absent_in_source" if synthetic_id else "unresolved"
            ),
            "title_raw": caption,
            "caption": caption,
            "footnotes": "\n".join(
                dict.fromkeys(part["footnotes"] for part in parts if part["footnotes"])
            ),
            "source_physical_pages": source_pages,
            "source_fragments": [
                {
                    key: part[key]
                    for key in (
                        "fragment_id",
                        "segment_id",
                        "block_index",
                        "physical_pages",
                        "bbox",
                        "image_path",
                        "standalone_title",
                        "association_reason",
                    )
                }
                for part in parts
            ],
            "components": components,
            "grid": grid,
            "fragment_column_counts": fragment_column_counts,
        }
        issues = quality_issues(candidate)
        if (
            candidate["node_id_match_key"]
            and identifier_group_counts[candidate["node_id_match_key"]] > 1
        ):
            issues.append(
                _issue(
                    "duplicate_identifier_candidates",
                    "More than one logical candidate normalizes to this Table identifier; candidates remain separate.",
                    "review",
                )
            )
        catalog_matches = catalog_by_key.get(candidate["node_id_match_key"], [])
        exact_catalog_matches = [
            node
            for node in catalog_matches
            if node["node_id_raw"].casefold() == (candidate["node_id_raw"] or "").casefold()
        ]
        if len(catalog_matches) > 1 and len(exact_catalog_matches) != 1:
            issues.append(
                _issue(
                    "catalog_identifier_collision",
                    "Multiple catalog nodes normalize to this identifier; no silent merge was made.",
                    "review",
                    catalog_node_ids=[node["node_id_raw"] for node in catalog_matches],
                )
            )
        needs_review = any(issue["severity"] in {"review", "error"} for issue in issues)
        candidate["issues"] = issues
        candidate["data_status"] = "review_required" if needs_review else "machine_validated"
        candidate["validation_route"] = (
            "mineru_plus_rules_pending_luna"
            if needs_review
            else "mineru_plus_rules"
        )
        key = _safe_key(node_id or synthetic_id or group_key)
        if (
            candidate["node_id_match_key"]
            and identifier_group_counts[candidate["node_id_match_key"]] > 1
        ):
            key += "__" + f"part_{len(candidates) + 1:03d}"
        artifact_dir = run_dir / f"Table_{key}"
        if artifact_dir.exists() and not force:
            raise FileExistsError(f"Table artifact already exists: {artifact_dir}; use --force")
        if (artifact_dir / "human_review.json").is_file():
            raise ValueError("This Table has human edits; reconcile explicitly before reprocessing")
        artifact_dir.mkdir(parents=True, exist_ok=True)
        candidate["artifact_directory"] = artifact_dir.relative_to(run_dir).as_posix()
        _write_json(artifact_dir / "candidate.json", candidate)
        mineru_runs = []
        for segment_id in sorted({part["segment_id"] for part in parts}):
            summary_path = run_dir / "ocr" / segment_id / "result_summary.json"
            if summary_path.is_file():
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
                summary.pop("zip_url", None)
                mineru_runs.append(summary)
        segment_records = [
            segment
            for segment in segments
            if segment["segment_id"] in {part["segment_id"] for part in parts}
        ]
        _write_json(
            artifact_dir / "table_provenance.json",
            {
                "schema_version": 1,
                "manual_pdf": str(pdf_path),
                "source_physical_pages": source_pages,
                "segments": segment_records,
                "mineru_runs": mineru_runs,
                "validation_route": candidate["validation_route"],
                "machine_issues": issues,
                "luna_patch_applied": False,
                "unresolved_issues": [
                    issue for issue in issues if issue["severity"] in {"review", "error"}
                ],
            },
        )
        (artifact_dir / "candidate.html").write_text(combined_html, encoding="utf-8")
        (artifact_dir / "notes.md").write_text(_notes(candidate), encoding="utf-8")
        _render_candidate_preview(candidate, artifact_dir / "candidate_preview.png")
        _write_workbook(artifact_dir / "table.xlsx", candidate)
        if needs_review:
            _write_json(artifact_dir / "review_request.json", _review_request(candidate))
        candidates.append(candidate)

    write_index(run_dir, candidates)
    merged = _merge_manifest_tables(manual_manifest, candidates)
    _write_json(run_dir / "manifest.json", merged)
    summary = {
        "mineru_table_fragments": len(fragments),
        "logical_tables": len(candidates),
        "machine_validated": sum(
            candidate["data_status"] == "machine_validated" for candidate in candidates
        ),
        "review_required": sum(
            candidate["data_status"] == "review_required" for candidate in candidates
        ),
        "luna_invoked": False,
    }
    _write_json(run_dir / "processing_summary.json", summary)
    return summary
