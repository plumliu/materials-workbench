"""Sparse, auditable human corrections to final Table workbooks."""

from __future__ import annotations

import json
import os
from pathlib import Path
import shutil
import tempfile
from typing import Any

from openpyxl import load_workbook

from .table_processing import _excel_value


def review_path(run_dir: Path, artifact_dir: Path) -> Path:
    relative = artifact_dir.resolve().relative_to(run_dir.resolve())
    if len(relative.parts) != 1 or not relative.name.startswith("Table_"):
        raise ValueError("Table artifact must be directly under the tables directory")
    return artifact_dir / "human_review.json"


def sheets(candidate: dict[str, Any]) -> list[dict[str, Any]]:
    return candidate.get("components") or [
        {"component_id": "component_001", "grid": candidate["grid"],
         "source_physical_pages": candidate["source_physical_pages"]}
    ]


def _base_revision(candidate: dict[str, Any]) -> int:
    return candidate["revision"]


def _covered_cells(grid: dict[str, Any]) -> set[tuple[int, int]]:
    covered: set[tuple[int, int]] = set()
    for cell in grid.get("cells", []):
        row, column = int(cell["row"]), int(cell["column"])
        for r in range(row, row + int(cell.get("rowspan", 1))):
            for c in range(column, column + int(cell.get("colspan", 1))):
                if (r, c) != (row, column):
                    covered.add((r, c))
    return covered


def validate_review(candidate: dict[str, Any], review: dict[str, Any]) -> None:
    if review.get("schema_version") != 2 or review.get("candidate_id") != candidate["candidate_id"]:
        raise ValueError("Human review belongs to a different candidate")
    if review.get("base_revision") != _base_revision(candidate):
        raise ValueError("Table changed since this review; review it again")
    if review.get("status") not in {"draft", "verified"}:
        raise ValueError("Invalid human review status")
    if not isinstance(review.get("notes"), str):
        raise ValueError("Human review notes must be text")
    pages = candidate["source_physical_pages"]
    checked = review.get("reviewed_pages")
    if not isinstance(checked, list) or any(type(page) is not int for page in checked):
        raise ValueError("Invalid reviewed pages")
    if len(checked) != len(set(checked)) or not set(checked) <= set(pages):
        raise ValueError("Reviewed pages must belong to this Table")
    if review["status"] == "verified" and set(checked) != set(pages):
        raise ValueError("Check every source page before confirming")
    changes = review.get("changes")
    if not isinstance(changes, list):
        raise ValueError("Changes must be a list")
    by_id = {part["component_id"]: part["grid"] for part in sheets(candidate)}
    covered = {key: _covered_cells(grid) for key, grid in by_id.items()}
    seen: set[tuple[str, int, int]] = set()
    for change in changes:
        if not isinstance(change, dict):
            raise ValueError("Invalid cell change")
        component = change.get("component_id")
        row, column = change.get("row"), change.get("column")
        if component not in by_id or type(row) is not int or type(column) is not int:
            raise ValueError("Invalid cell address")
        rows = by_id[component]["rows"]
        if row < 1 or row > len(rows) or column < 1 or column > len(rows[row - 1]):
            raise ValueError("Cell address is outside the Table")
        key = (component, row, column)
        if key in seen or (row, column) in covered[component]:
            raise ValueError("Duplicate or merged placeholder cell")
        seen.add(key)
        if not isinstance(change.get("before"), str) or not isinstance(change.get("after"), str):
            raise ValueError("Cell values must be text")
        if rows[row - 1][column - 1] != change["before"] or change["after"] == change["before"]:
            raise ValueError("Cell base value changed or correction is empty")


def read_review(path: Path, candidate: dict[str, Any]) -> dict[str, Any]:
    if not path.is_file():
        return {
            "schema_version": 2, "candidate_id": candidate["candidate_id"],
            "base_revision": _base_revision(candidate),
            "revision": 0, "status": "draft", "reviewed_pages": [],
            "notes": "", "changes": [],
        }
    review = json.loads(path.read_text(encoding="utf-8"))
    validate_review(candidate, review)
    return review


def save_review(path: Path, candidate: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    current = editable_review(path, candidate)
    if incoming.get("base_revision") != current["base_revision"]:
        raise ValueError("Table changed since it was opened; reload before saving")
    if incoming.get("revision") != current["revision"]:
        raise ValueError("Review changed in another window; reload before saving")
    review = {
        "schema_version": 2,
        "candidate_id": candidate["candidate_id"],
        "base_revision": _base_revision(candidate),
        "revision": current["revision"] + 1,
        "status": incoming.get("status"),
        "reviewed_pages": incoming.get("reviewed_pages"),
        "notes": incoming.get("notes"),
        "changes": incoming.get("changes"),
    }
    validate_review(candidate, review)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        json.dump(review, handle, ensure_ascii=False, indent=2)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return review


def editable_review(path: Path, candidate: dict[str, Any]) -> dict[str, Any]:
    """Keep stale edits on disk until the user explicitly saves a fresh review."""
    try:
        return read_review(path, candidate)
    except ValueError:
        previous = json.loads(path.read_text(encoding="utf-8"))
        if previous.get("schema_version") != 2 or previous.get("candidate_id") != candidate["candidate_id"] or previous.get("base_revision") == _base_revision(candidate):
            raise
        return {"schema_version": 2, "candidate_id": candidate["candidate_id"],
                "base_revision": _base_revision(candidate), "revision": previous["revision"],
                "status": "draft", "reviewed_pages": [], "notes": "", "changes": [], "stale": True}


def apply_verified_review(artifact_dir: Path, output_dir: Path, run_dir: Path) -> bool:
    path = review_path(run_dir, artifact_dir)
    if not path.is_file():
        return False
    if json.loads(path.read_text(encoding="utf-8")).get("status") != "verified":
        return False
    candidate = json.loads((artifact_dir / "candidate.json").read_text(encoding="utf-8"))
    review = read_review(path, candidate)
    workbook = load_workbook(output_dir / "table.xlsx")
    for change in review["changes"]:
        index = next(
            i for i, part in enumerate(sheets(candidate))
            if part["component_id"] == change["component_id"]
        )
        cell = workbook.worksheets[index].cell(change["row"], change["column"])
        cell.value = _excel_value(change["after"])
        if change["after"].startswith("="):
            cell.data_type = "s"
    workbook.save(output_dir / "table.xlsx")
    shutil.copy2(path, output_dir / "human_review.json")
    provenance_path = output_dir / "table_provenance.json"
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    provenance["human_review"] = {
        "status": "verified", "reviewed_pages": review["reviewed_pages"],
        "changed_cells": len(review["changes"]), "record": "human_review.json",
    }
    provenance_path.write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    return True
