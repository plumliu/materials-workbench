import json
import shutil
from http.server import ThreadingHTTPServer
from threading import Thread
from urllib.parse import urlencode
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
import pymupdf
import pytest

from pdf_tree_workflow.human_review import apply_verified_review, read_review, review_path, save_review
from pdf_tree_workflow.review_server import ReviewHandler


def test_verified_review_applies_only_to_assembled_copy(tmp_path):
    artifact = tmp_path / "Table_1.2"
    artifact.mkdir(parents=True)
    candidate = {
        "schema_version": 2, "revision": 1, "candidate_id": "table-1", "source_physical_pages": [3, 4],
        "grid": {"rows": [["old"]], "cells": []},
    }
    (artifact / "candidate.json").write_text(json.dumps(candidate), encoding="utf-8")
    workbook = Workbook()
    workbook.active["A1"] = "old"
    workbook.save(artifact / "table.xlsx")
    (artifact / "table_provenance.json").write_text("{}", encoding="utf-8")
    review_file = review_path(tmp_path, artifact)
    change = {"component_id": "component_001", "row": 1, "column": 1, "before": "old", "after": "5"}
    draft = save_review(review_file, candidate, {
        "base_revision": read_review(review_file, candidate)["base_revision"],
        "revision": 0, "status": "draft", "reviewed_pages": [3],
        "notes": "Check continuation page", "changes": [change],
    })
    output = tmp_path / "library_table"
    output.mkdir()
    shutil.copy2(artifact / "table.xlsx", output / "table.xlsx")
    shutil.copy2(artifact / "table_provenance.json", output / "table_provenance.json")
    assert not apply_verified_review(artifact, output, tmp_path)
    with pytest.raises(ValueError, match="every source page"):
        save_review(review_file, candidate, {
            **draft, "status": "verified", "reviewed_pages": [3],
        })
    verified = save_review(review_file, candidate, {
        **draft, "status": "verified", "reviewed_pages": [3, 4],
    })
    with pytest.raises(ValueError, match="another window"):
        save_review(review_file, candidate, {**draft, "status": "draft"})
    assert verified["revision"] == 2
    assert apply_verified_review(artifact, output, tmp_path)
    final = load_workbook(output / "table.xlsx", read_only=True)
    try:
        assert final.active["A1"].value == 5
    finally:
        final.close()
    source = load_workbook(artifact / "table.xlsx", read_only=True)
    try:
        assert source.active["A1"].value == "old"
    finally:
        source.close()
    assert json.loads((output / "table_provenance.json").read_text())["human_review"]["changed_cells"] == 1
    assert (output / "human_review.json").is_file()
    candidate["grid"]["rows"][0][0] = "new base"
    candidate["revision"] = 2
    with pytest.raises(ValueError, match="Table changed"):
        read_review(review_file, candidate)


def test_two_page_table_serves_both_original_pages(tmp_path):
    run = tmp_path / "manual" / "tables"
    artifact = run / "Table_1"
    pages = tmp_path / "sources"
    pages.mkdir(parents=True)
    artifact.mkdir(parents=True)
    (run / "manifest.json").write_text(json.dumps({"page_review_dir": str(pages)}))
    candidate = {
        "schema_version": 2, "revision": 1, "candidate_id": "table-1", "caption": "Table 1", "source_physical_pages": [3, 4],
        "artifact_directory": "Table_1", "data_status": "machine_validated",
        "grid": {"rows": [["value"]], "cells": []},
    }
    (artifact / "candidate.json").write_text(json.dumps(candidate), encoding="utf-8")
    (run / "index.json").write_text(json.dumps([candidate]), encoding="utf-8")
    for page in (3, 4):
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), f"physical page {page}")
            document.save(pages / f"page_{page:04d}.pdf")
    ReviewHandler.runs_dir = tmp_path
    server = ThreadingHTTPServer(("127.0.0.1", 0), ReviewHandler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        query = urlencode({"manual": "manual", "id": "table-1"})
        with urlopen(f"http://127.0.0.1:{server.server_port}/api/table?{query}") as response:
            assert json.load(response)["pages"] == [3, 4]
        for page in (3, 4):
            with urlopen(f"http://127.0.0.1:{server.server_port}/api/page-image?{query}&page={page}") as response:
                assert response.read(8) == b"\x89PNG\r\n\x1a\n"
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
